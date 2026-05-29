from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
import re
import io
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile, ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pypdf import PdfReader


@dataclass(frozen=True)
class FormatFinding:
    rule_id: str
    priority: str
    location: str
    message: str
    suggestion: str


@dataclass(frozen=True)
class ParsedDocument:
    file_name: str
    text: str
    markdown: str
    parser_name: str
    format_findings: tuple[str, ...] = ()
    format_rule_findings: tuple[FormatFinding, ...] = ()
    tables: tuple[tuple[tuple[str, ...], ...], ...] = ()


def parse_document(file_name: str, content: bytes, parsing_instruction: str | None = None) -> ParsedDocument:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(file_name, content, parsing_instruction)
    if suffix == ".docx":
        return parse_docx(file_name, content)
    if suffix == ".doc":
        return parse_doc(file_name, content)
    raise ValueError(f"不支持的文件类型：{suffix}")


def parse_pdf(file_name: str, content: bytes, parsing_instruction: str | None = None) -> ParsedDocument:
    api_key = os.getenv("LLAMA_CLOUD_API_KEY")
    if api_key:
        try:
            parsed = _parse_with_llamaparse(file_name, content, api_key, parsing_instruction)
            if parsed.text.strip():
                return parsed
            fallback = _parse_with_pypdf(file_name, content)
            return ParsedDocument(
                file_name=file_name,
                text=fallback.text,
                markdown=fallback.markdown + "\n\n> LlamaParse 未返回有效文本，已使用本地文本解析。",
                parser_name="pypdf-fallback-empty-llamaparse",
                format_rule_findings=fallback.format_rule_findings,
                tables=fallback.tables,
            )
        except Exception as exc:  # pragma: no cover - Streamlit displays fallback result.
            fallback = _parse_with_pypdf(file_name, content)
            return ParsedDocument(
                file_name=file_name,
                text=fallback.text,
                markdown=fallback.markdown + f"\n\n> LlamaParse 解析失败，已使用本地文本解析。错误：{exc}",
                parser_name="pypdf-fallback",
                format_rule_findings=fallback.format_rule_findings,
                tables=fallback.tables,
            )
    return _parse_with_pypdf(file_name, content)


def _parse_with_llamaparse(file_name: str, content: bytes, api_key: str, parsing_instruction: str | None) -> ParsedDocument:
    from llama_parse import LlamaParse

    with tempfile.NamedTemporaryFile(suffix=Path(file_name).suffix or ".pdf") as temp:
        temp.write(content)
        temp.flush()
        parser = LlamaParse(
            api_key=api_key,
            result_type="markdown",
            language="chinese",
            premium_mode=True,
            adaptive_long_table=True,
            aggressive_table_extraction=True,
            merge_tables_across_pages_in_markdown=True,
            preserve_layout_alignment_across_pages=True,
            parsing_instruction=parsing_instruction or "请保留中文教学大纲中的复杂表格结构、合并单元格关系、标题层级和所有单元格内容。",
        )
        documents = parser.load_data(temp.name)
    markdown = "\n\n".join(doc.text for doc in documents)
    return ParsedDocument(file_name=file_name, text=markdown, markdown=markdown, parser_name="LlamaParse", tables=_extract_markdown_tables(markdown))


def _parse_with_pypdf(file_name: str, content: bytes) -> ParsedDocument:
    with tempfile.NamedTemporaryFile(suffix=".pdf") as temp:
        temp.write(content)
        temp.flush()
        reader = PdfReader(temp.name)
        pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(pages)
    return ParsedDocument(file_name=file_name, text=text, markdown=text, parser_name="pypdf", tables=_extract_markdown_tables(text))


def parse_docx(file_name: str, content: bytes) -> ParsedDocument:
    local = _parse_docx_locally(file_name, content)
    api_key = os.getenv("LLAMA_CLOUD_API_KEY")
    if not api_key:
        return local
    try:
        parsed = _parse_with_llamaparse(file_name, content, api_key, "请解析 Word 版中文教学大纲，保留标题层级、复杂表格结构、合并单元格关系和所有单元格内容。")
    except Exception:
        return local
    if not parsed.text.strip():
        return local
    return ParsedDocument(
        file_name=file_name,
        text=parsed.text,
        markdown=parsed.markdown,
        parser_name="LlamaParse + python-docx",
        format_findings=local.format_findings,
        format_rule_findings=local.format_rule_findings,
        tables=_merge_tables(parsed.tables, local.tables),
    )


def _parse_docx_locally(file_name: str, content: bytes) -> ParsedDocument:
    parser_name = "python-docx"
    with tempfile.NamedTemporaryFile(suffix=".docx") as temp:
        temp.write(content)
        temp.flush()
        try:
            document = Document(temp.name)
        except BadZipFile:
            repaired_content = _repair_docx_without_bad_parts(content)
            if repaired_content is None:
                raise
            temp.seek(0)
            temp.truncate()
            temp.write(repaired_content)
            temp.flush()
            document = Document(temp.name)
            parser_name = "python-docx-repaired"
        format_findings = _inspect_docx_cleanliness(temp.name)
        format_rule_findings = _inspect_docx_format(document)

    blocks: list[str] = []
    tables: list[tuple[tuple[str, ...], ...]] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)

    for table in document.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
        if rows:
            tables.append(tuple(tuple(cell for cell in row) for row in rows))
            blocks.append(_table_to_markdown(rows))

    markdown = "\n\n".join(blocks)
    return ParsedDocument(
        file_name=file_name,
        text=markdown,
        markdown=markdown,
        parser_name=parser_name,
        format_findings=tuple(format_findings),
        format_rule_findings=tuple(format_rule_findings),
        tables=tuple(tables),
    )


def _repair_docx_without_bad_parts(content: bytes) -> bytes | None:
    output = io.BytesIO()
    skipped: set[str] = set()
    copied: list[tuple[str, bytes]] = []
    try:
        with ZipFile(io.BytesIO(content)) as source:
            for info in source.infolist():
                try:
                    copied.append((info.filename, source.read(info.filename)))
                except BadZipFile:
                    skipped.add(info.filename)
    except BadZipFile:
        return None

    if not skipped:
        return None

    with ZipFile(output, "w", ZIP_DEFLATED) as target:
        for filename, data in copied:
            if filename in skipped:
                continue
            if filename.endswith(".rels"):
                data = _remove_rels_to_skipped_parts(filename, data, skipped)
            target.writestr(filename, data)
    return output.getvalue()


def _remove_rels_to_skipped_parts(rels_filename: str, data: bytes, skipped: set[str]) -> bytes:
    try:
        xml = data.decode("utf-8")
    except UnicodeDecodeError:
        return data
    base_dir = str(Path(rels_filename).parent.parent)
    if base_dir == ".":
        base_dir = ""

    def should_remove(match: re.Match[str]) -> bool:
        relationship = match.group(0)
        target_match = re.search(r'Target="([^"]+)"', relationship)
        if not target_match:
            return False
        target = target_match.group(1)
        if re.match(r"^[a-z]+:", target) or target.startswith("/"):
            return False
        normalized = str(Path(base_dir, target)).replace("\\", "/")
        return normalized in skipped

    xml = re.sub(r"<Relationship\b[^>]*/>", lambda match: "" if should_remove(match) else match.group(0), xml)
    return xml.encode("utf-8")


def _merge_tables(
    primary: tuple[tuple[tuple[str, ...], ...], ...],
    secondary: tuple[tuple[tuple[str, ...], ...], ...],
) -> tuple[tuple[tuple[str, ...], ...], ...]:
    seen: set[tuple[tuple[str, ...], ...]] = set()
    merged: list[tuple[tuple[str, ...], ...]] = []
    for table in primary + secondary:
        if table not in seen:
            seen.add(table)
            merged.append(table)
    return tuple(merged)


def parse_doc(file_name: str, content: bytes) -> ParsedDocument:
    api_key = os.getenv("LLAMA_CLOUD_API_KEY")
    if api_key:
        try:
            return _parse_with_llamaparse(file_name, content, api_key, "请解析 Word 版中文教学大纲，保留标题层级、表格结构和所有单元格内容。")
        except Exception:
            pass
    if not shutil.which("textutil"):
        raise ValueError(".doc 文件需要 LlamaParse API 或本机 textutil/LibreOffice 转换支持。建议优先上传 .docx 或 PDF。")
    with tempfile.TemporaryDirectory() as temp_dir:
        source = Path(temp_dir) / file_name
        source.write_bytes(content)
        output = Path(temp_dir) / f"{source.stem}.txt"
        command = ["textutil", "-convert", "txt", "-output", str(output), str(source)]
        completed = subprocess.run(command, capture_output=True, text=True, check=False)
        if completed.returncode != 0 or not output.exists():
            message = completed.stderr.strip() or "系统未能将 .doc 文件转换为文本。"
            raise ValueError(f".doc 文件解析失败：{message}")
        text = output.read_text(encoding="utf-8", errors="ignore")
    return ParsedDocument(file_name=file_name, text=text, markdown=text, parser_name="textutil-doc")


def _table_to_markdown(rows: list[list[str]]) -> str:
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    body = normalized[1:] or [[""] * width]
    lines = [
        "| " + " | ".join(_escape_table_cell(cell) for cell in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(_escape_table_cell(cell) for cell in row) + " |")
    return "\n".join(lines)


def _escape_table_cell(value: str) -> str:
    return value.replace("|", "\\|")


def _inspect_docx_format(document: Document) -> list[FormatFinding]:
    findings: list[FormatFinding] = []
    findings.extend(_inspect_page_margins(document))
    findings.extend(_inspect_body_paragraph_format(document))
    findings.extend(_inspect_course_target_intro_bold(document))
    findings.extend(_inspect_heading_format(document))
    findings.extend(_inspect_table_titles(document))
    findings.extend(_inspect_table_format(document))
    findings.extend(_inspect_table_cell_alignment(document))
    return findings


def _inspect_page_margins(document: Document) -> list[FormatFinding]:
    bad_sections = []
    for index, section in enumerate(document.sections, start=1):
        margins = {
            "上": _cm(section.top_margin),
            "下": _cm(section.bottom_margin),
            "左": _cm(section.left_margin),
            "右": _cm(section.right_margin),
        }
        bad = [f"{name}{value:.2f}cm" for name, value in margins.items() if value is not None and abs(value - 2.0) > 0.08]
        if bad:
            bad_sections.append(f"第{index}节（{', '.join(bad)}）")
    if not bad_sections:
        return []
    return [
        FormatFinding(
            "R-FORMAT-101",
            "中",
            "页面设置",
            "页边距不符合上下左右均为 2cm 的要求：" + "；".join(bad_sections[:3]) + "。",
            "请将页面上、下、左、右边距统一设置为 2cm。",
        )
    ]


def _inspect_body_paragraph_format(document: Document) -> list[FormatFinding]:
    font_examples: list[str] = []
    size_examples: list[str] = []
    line_examples: list[str] = []
    indent_examples: list[str] = []
    for paragraph in document.paragraphs:
        text = _compact_text(paragraph.text)
        if not _looks_like_body_paragraph(text):
            continue
        if not _paragraph_uses_songti(paragraph, document):
            font_examples.append(text)
        if not _paragraph_size_ok(paragraph, document, 12.0, tolerance=0.6):
            size_examples.append(text)
        if not _paragraph_line_spacing_ok(paragraph, 22.0, tolerance=1.2):
            line_examples.append(text)
        if not _paragraph_has_numbering(paragraph) and not _first_line_indent_ok(paragraph, min_cm=0.65, max_cm=1.05):
            indent_examples.append(text)

    findings: list[FormatFinding] = []
    if font_examples:
        findings.append(
            FormatFinding(
                "R-FORMAT-102",
                "低",
                "正文格式",
                _format_examples("正文疑似未使用宋体字体", font_examples),
                "请将正文中文字体统一设置为宋体。",
            )
        )
    if size_examples:
        findings.append(
            FormatFinding(
                "R-FORMAT-103",
                "低",
                "正文格式",
                _format_examples("正文疑似未使用小四号（约 12pt）", size_examples),
                "请将正文统一设置为小四号（约 12pt）。",
            )
        )
    if line_examples:
        findings.append(
            FormatFinding(
                "R-FORMAT-104",
                "低",
                "正文格式",
                _format_examples("正文行距疑似不是 22 磅", line_examples),
                "请将正文行距设置为固定值或最小值 22 磅。",
            )
        )
    if indent_examples:
        findings.append(
            FormatFinding(
                "R-FORMAT-105",
                "低",
                "正文格式",
                _format_examples("正文首行缩进疑似不是约 2 字符", indent_examples),
                "请将正文段落首行缩进设置为约 2 字符。",
            )
        )
    return findings


def _inspect_course_target_intro_bold(document: Document) -> list[FormatFinding]:
    bad_targets: list[str] = []
    for paragraph in document.paragraphs:
        text = _compact_text(paragraph.text)
        if "评分标准" in text:
            continue
        match = re.match(r"^(课程目标\s*[一二三四五六七八九十\d]+)\s*[:：]", text)
        if not match:
            continue
        if not _paragraph_prefix_bold(paragraph, match.group(1)):
            bad_targets.append(match.group(1) + "：")
    if not bad_targets:
        return []
    return [
        FormatFinding(
            "R-FORMAT-112",
            "低",
            "课程目标格式",
            "“（一）课程目标”下课程目标段落开头未加粗，例如：" + "、".join(bad_targets[:5]) + "。",
            "请将每段课程目标开头的“课程目标1：”“课程目标2：”等文字加粗。",
        )
    ]


def _paragraph_prefix_bold(paragraph, prefix: str) -> bool:
    remaining = re.sub(r"\s+", "", prefix)
    for run in paragraph.runs:
        text = re.sub(r"\s+", "", run.text or "")
        if not text:
            continue
        take = min(len(text), len(remaining))
        if take <= 0:
            return True
        if text[:take] and not (run.bold is True or _run_has_bold(run)):
            return False
        remaining = remaining[take:]
        if not remaining:
            return True
    return False


def _inspect_heading_format(document: Document) -> list[FormatFinding]:
    level1_bad: list[str] = []
    level2_bad: list[str] = []
    for paragraph in document.paragraphs:
        text = _compact_text(paragraph.text)
        if not text:
            continue
        if re.match(r"^[一二三四五六七八九十]+、", text):
            if not (_paragraph_uses_heiti(paragraph, document) and _paragraph_size_ok(paragraph, document, 14.0, tolerance=1.1)):
                level1_bad.append(text)
        elif re.match(r"^（[一二三四五六七八九十]+）", text):
            if not (_paragraph_uses_heiti(paragraph, document) and _paragraph_size_ok(paragraph, document, 12.0, tolerance=0.8)):
                level2_bad.append(text)

    findings: list[FormatFinding] = []
    if level1_bad:
        findings.append(
            FormatFinding(
                "R-FORMAT-106",
                "低",
                "标题格式",
                _format_examples("一级标题格式不符合“一、二、三”编号、黑体、四号（约 14pt）要求", level1_bad),
                "请将一级标题设置为“一、二、三”编号格式，字体为黑体，字号约 14pt（四号）。",
            )
        )
    if level2_bad:
        findings.append(
            FormatFinding(
                "R-FORMAT-107",
                "低",
                "标题格式",
                _format_examples("二级标题格式不符合“（一）（二）”编号、黑体、小四（约 12pt）要求", level2_bad),
                "请将二级标题设置为“（一）（二）”编号格式，字体为黑体，字号约 12pt（小四）。",
            )
        )
    return findings


def _inspect_table_titles(document: Document) -> list[FormatFinding]:
    bad_titles: list[str] = []
    for paragraph in document.paragraphs:
        text = _compact_text(paragraph.text)
        if not text:
            continue
        if text.startswith("表"):
            if not re.match(r"^表\s*\d+(?:-\d+)?\s*\S+", text):
                bad_titles.append(text)
                continue
            if paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER or not _paragraph_has_bold(paragraph):
                bad_titles.append(text)
    if not bad_titles:
        return []
    return [
        FormatFinding(
            "R-FORMAT-108",
            "低",
            "表题格式",
            _format_examples("表题格式不符合居中、加粗、以“表+编号+名称”开头的要求", bad_titles),
            "请将表题设置为居中、加粗，并使用“表+编号+名称”的规范写法。",
        )
    ]


def _inspect_table_format(document: Document) -> list[FormatFinding]:
    font_size_cells: list[str] = []
    table_titles = _table_titles(document)
    for index, table in enumerate(document.tables, start=1):
        font_size_cells.extend(_table_font_size_findings(table, table_titles.get(index, f"第{index}个表格")))

    findings: list[FormatFinding] = []
    if font_size_cells:
        findings.append(
            FormatFinding(
                "R-FORMAT-110",
                "低",
                "表格格式",
                "表格内字体字号疑似不符合五号（约 10.5pt）或模板允许的小字号要求：" + "；".join(font_size_cells[:6]) + "。",
                "请将表格内文字一般设置为五号（约 10.5pt）；内容较密的大表可按模板适当使用约 9pt，但应保持全文一致。",
            )
        )
    return findings


def _table_titles(document: Document) -> dict[int, str]:
    titles: dict[int, str] = {}
    table_index = 0
    for block in document.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = _normalize_repeated_text(_compact_text("".join(node.text or "" for node in block.findall(".//" + _w("t")))))
            if re.match(r"^表\s*\d+(?:-\d+)?\s*\S+", text):
                titles[table_index + 1] = text
        elif tag == "tbl":
            table_index += 1
    return titles


def _table_font_size_findings(table, table_name: str) -> list[str]:
    findings: list[str] = []
    header_rows = _table_header_rows(table)
    column_names = _table_column_names(header_rows)
    seen: set[str] = set()
    for row_index, row in enumerate(table.rows, start=1):
        for col_index, cell in enumerate(row.cells):
            bad_sizes: list[float] = []
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if not run.text.strip():
                        continue
                    size = _run_size_pt(run)
                    if size is not None and not 8.8 <= size <= 10.8:
                        bad_sizes.append(size)
            if not bad_sizes:
                continue
            column_name = column_names.get(col_index, f"第{col_index + 1}列")
            key = f"{table_name}-{column_name}"
            if key in seen:
                continue
            seen.add(key)
            size_text = "、".join(f"{size:g}pt" for size in sorted(set(round(size, 1) for size in bad_sizes)))
            findings.append(f"{table_name}的“{column_name}”列字号为{size_text}")
    return findings


def _table_header_rows(table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows[:2]:
        rows.append([_compact_text(cell.text) for cell in row.cells])
    return rows


def _table_column_names(header_rows: list[list[str]]) -> dict[int, str]:
    if not header_rows:
        return {}
    first_row = [_normalize_repeated_text(value) for value in header_rows[0]]
    duplicate_headers = {value for value in first_row if value and first_row.count(value) > 1}
    generic_headers = {"学时安排", "平时考核"}
    names: dict[int, str] = {}
    for index, value in enumerate(first_row):
        if value:
            names[index] = value
    for row in header_rows[1:]:
        for index, raw_value in enumerate(row):
            value = _normalize_repeated_text(raw_value)
            if not value:
                continue
            current = names.get(index, "")
            if not current:
                names[index] = value
            elif (current in duplicate_headers or current in generic_headers) and value != current and value not in current:
                names[index] = f"{current}-{value}"
    return names


def _normalize_repeated_text(text: str) -> str:
    text = _compact_text(text)
    if not text:
        return ""
    half = len(text) // 2
    if half and len(text) % 2 == 0 and text[:half] == text[half:]:
        return text[:half]
    for length in range(1, len(text) // 2 + 1):
        if len(text) % length == 0:
            part = text[:length]
            if part and part * (len(text) // length) == text:
                return part
    return text


def _inspect_table_cell_alignment(document: Document) -> list[FormatFinding]:
    bad_cells: list[str] = []
    for table_index, table in enumerate(document.tables, start=1):
        header_map = _table_header_indices(table)
        if not {"课程目标", "支撑的毕业要求"}.issubset(header_map):
            continue
        target_columns = [header_map["课程目标"], header_map["支撑的毕业要求"]]
        for row_index, row in enumerate(table.rows[1:], start=2):
            for column_index in target_columns:
                if column_index >= len(row.cells):
                    continue
                cell = row.cells[column_index]
                text = _compact_text(cell.text)
                if not text or text in {"课程目标", "支撑的毕业要求"}:
                    continue
                if cell.vertical_alignment != WD_CELL_VERTICAL_ALIGNMENT.CENTER:
                    bad_cells.append(f"第{row_index}行“{text[:18]}”")
    if not bad_cells:
        return []
    return [
        FormatFinding(
            "R-FORMAT-111",
            "低",
            "表格格式",
            "课程目标与毕业要求关系表中部分单元格未设置垂直居中：" + "、".join(bad_cells[:8]) + "。",
            "请将“课程目标”和“支撑的毕业要求”列中的内容统一设置为水平、垂直居中。",
        )
    ]


def _table_header_indices(table) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for row in table.rows[:3]:
        for column_index, cell in enumerate(row.cells):
            clean = _compact_text(cell.text)
            if clean in {"课程目标", "支撑的毕业要求"}:
                header_map[clean] = column_index
    return header_map


def _looks_like_body_paragraph(text: str) -> bool:
    if len(text) < 12:
        return False
    if re.search(r"https?://|www\.", text):
        return False
    if re.match(r"^(注|备注)\s*[:：]", text):
        return False
    if re.match(r"^（\d+）", text) and re.search(r"\[[A-Z]\]|出版社|出版|20\d{2}", text):
        return False
    if re.match(r"^课程目标[\d一二三四五六七八九十、,，和及\s]+的?评分标准$", text):
        return False
    if re.match(r"^[一二三四五六七八九十]+、|^（[一二三四五六七八九十]+）|^\d+\.|^表\s*\d+", text):
        return False
    if text.startswith("西京学院《") or "教学大纲（2026版）" in text:
        return False
    if re.match(r"^(知识模块|实践项目|模块[一二三四五六七八九十]+)", text):
        return False
    if re.search(r"课程基本信息|课程目标与毕业要求的关系|教学内容与课程目标的支撑关系|课程目标考核环节评定方式", text):
        return False
    return True


def _paragraph_uses_songti(paragraph, document: Document) -> bool:
    fonts = _paragraph_direct_east_asia_fonts(paragraph)
    if fonts:
        return all(_font_is_songti(font) for font in fonts)
    inherited_font = _style_east_asia_font(paragraph.style) or _normal_style_east_asia_font(document)
    return True if not inherited_font else _font_is_songti(inherited_font)


def _paragraph_uses_heiti(paragraph, document: Document) -> bool:
    fonts = _paragraph_direct_east_asia_fonts(paragraph)
    if fonts:
        return all(_font_is_heiti(font) for font in fonts)
    return _font_is_heiti(_style_east_asia_font(paragraph.style) or _normal_style_east_asia_font(document))


def _paragraph_direct_east_asia_fonts(paragraph) -> list[str]:
    fonts: list[str] = []
    for run in paragraph.runs:
        if not run.text.strip() or not re.search(r"[\u4e00-\u9fff]", run.text):
            continue
        font = _run_east_asia_font(run)
        if font:
            fonts.append(font)
    return fonts


def _paragraph_size_ok(paragraph, document: Document, expected_pt: float, tolerance: float) -> bool:
    sizes = [_run_size_pt(run) for run in paragraph.runs if run.text.strip() and _run_size_pt(run) is not None]
    if not sizes:
        style_size = _style_size_pt(paragraph.style) or _normal_style_size_pt(document)
        return style_size is None or abs(style_size - expected_pt) <= tolerance
    return all(abs(size - expected_pt) <= tolerance for size in sizes)


def _paragraph_line_spacing_ok(paragraph, expected_pt: float, tolerance: float) -> bool:
    spacing = paragraph.paragraph_format.line_spacing
    if hasattr(spacing, "pt"):
        return abs(spacing.pt - expected_pt) <= tolerance
    return True


def _first_line_indent_ok(paragraph, min_cm: float, max_cm: float) -> bool:
    indent = paragraph.paragraph_format.first_line_indent
    if indent is None:
        return True
    value = _cm(indent)
    return value is None or min_cm <= value <= max_cm


def _paragraph_has_numbering(paragraph) -> bool:
    ppr = paragraph._p.pPr
    return bool(ppr is not None and ppr.numPr is not None)


def _paragraph_has_bold(paragraph) -> bool:
    runs = [run for run in paragraph.runs if run.text.strip() and re.search(r"[\w\u4e00-\u9fff]", run.text)]
    return bool(runs) and all(run.bold is True or _run_has_bold(run) for run in runs)


def _table_has_full_borders(table) -> bool:
    borders = _table_border_names(table)
    required = {"top", "left", "bottom", "right", "insideH", "insideV"}
    if required.issubset(borders):
        return True
    style = table.style
    if style is not None:
        return required.issubset(_style_table_border_names(style))
    return False


def _table_border_names(table) -> set[str]:
    tbl_borders = table._tbl.find(".//" + _w("tblBorders"))
    return _border_names(tbl_borders)


def _style_table_border_names(style) -> set[str]:
    tbl_borders = style.element.find(".//" + _w("tblBorders"))
    return _border_names(tbl_borders)


def _border_names(tbl_borders) -> set[str]:
    if tbl_borders is None:
        return set()
    names = set()
    for child in tbl_borders:
        value = child.get(qn("w:val"))
        if value and value not in {"nil", "none"}:
            names.add(child.tag.rsplit("}", 1)[-1])
    return names


def _table_font_size_ok(table) -> bool:
    sizes: list[float] = []
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if not run.text.strip():
                        continue
                    size = _run_size_pt(run)
                    if size is not None:
                        sizes.append(size)
    if not sizes:
        return True
    return all(8.8 <= size <= 10.8 for size in sizes)


def _table_first_text(table) -> str:
    for row in table.rows:
        for cell in row.cells:
            text = _compact_text(cell.text)
            if text:
                return text[:30]
    return ""


def _run_east_asia_font(run) -> str:
    rpr = run._element.rPr
    if rpr is None or rpr.rFonts is None:
        return ""
    return rpr.rFonts.get(qn("w:eastAsia")) or ""


def _style_east_asia_font(style) -> str:
    if style is None:
        return ""
    rpr = style.element.find(_w("rPr"))
    if rpr is None:
        return ""
    rfonts = rpr.find(_w("rFonts"))
    return rfonts.get(qn("w:eastAsia")) if rfonts is not None else ""


def _normal_style_east_asia_font(document: Document) -> str:
    try:
        return _style_east_asia_font(document.styles["Normal"])
    except Exception:
        return ""


def _run_size_pt(run) -> float | None:
    if run.font.size is not None:
        return float(run.font.size.pt)
    rpr = run._element.rPr
    if rpr is None:
        return None
    size = rpr.find(_w("sz"))
    if size is None or not size.get(qn("w:val")):
        return None
    return float(size.get(qn("w:val"))) / 2


def _style_size_pt(style) -> float | None:
    if style is None:
        return None
    if getattr(style, "font", None) is not None and style.font.size is not None:
        return float(style.font.size.pt)
    rpr = style.element.find(_w("rPr"))
    if rpr is None:
        return None
    size = rpr.find(_w("sz"))
    if size is None or not size.get(qn("w:val")):
        return None
    return float(size.get(qn("w:val"))) / 2


def _normal_style_size_pt(document: Document) -> float | None:
    try:
        return _style_size_pt(document.styles["Normal"])
    except Exception:
        return None


def _run_has_bold(run) -> bool:
    rpr = run._element.rPr
    return rpr is not None and rpr.find(_w("b")) is not None


def _font_is_songti(font: str) -> bool:
    return bool(font) and ("宋体" in font or "SimSun" in font)


def _font_is_heiti(font: str) -> bool:
    return bool(font) and ("黑体" in font or "SimHei" in font)


def _cm(value) -> float | None:
    return None if value is None else float(value.cm)


def _w(tag: str) -> str:
    return "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + tag


def _compact_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _format_examples(prefix: str, examples: list[str]) -> str:
    unique = []
    for example in examples:
        if example not in unique:
            unique.append(example)
    preview = "；".join(f"“{example[:40]}”" for example in unique[:3])
    extra = f"等 {len(unique)} 处" if len(unique) > 3 else f"{len(unique)} 处"
    return f"{prefix}，例如：{preview}，共{extra}。"


def _inspect_docx_cleanliness(path: str) -> list[str]:
    findings: list[str] = []
    try:
        with ZipFile(path) as archive:
            names = set(archive.namelist())
            if "word/comments.xml" in names:
                comments_xml = archive.read("word/comments.xml").decode("utf-8", errors="ignore")
                for comment_text in _extract_word_text_fragments(comments_xml):
                    if comment_text:
                        findings.append(f"存在批注：{comment_text[:80]}")
            if "word/document.xml" in names:
                document_xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
                for tag, label in [("ins", "插入修订"), ("del", "删除修订"), ("moveFrom", "移动前修订"), ("moveTo", "移动后修订")]:
                    for body in re.findall(rf"<w:{tag}\b[^>]*>(.*?)</w:{tag}>", document_xml, flags=re.DOTALL):
                        text = "".join(_extract_word_text_fragments(body)).strip()
                        findings.append(f"存在{label}" + (f"：{text[:80]}" if text else ""))
    except OSError:
        return findings
    return findings


def _extract_word_text_fragments(xml: str) -> list[str]:
    return [
        re.sub(r"\s+", " ", fragment).strip()
        for fragment in re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml, flags=re.DOTALL)
        if fragment.strip()
    ]


def _extract_markdown_tables(markdown: str) -> tuple[tuple[tuple[str, ...], ...], ...]:
    tables: list[tuple[tuple[str, ...], ...]] = []
    current: list[str] = []
    for line in markdown.splitlines():
        if line.strip().startswith("|") and "|" in line:
            current.append(line.strip())
            continue
        if current:
            parsed = _parse_markdown_table(current)
            if parsed:
                tables.append(parsed)
            current = []
    if current:
        parsed = _parse_markdown_table(current)
        if parsed:
            tables.append(parsed)
    return tuple(tables)


def _parse_markdown_table(lines: list[str]) -> tuple[tuple[str, ...], ...] | None:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(tuple(cells))
    return tuple(rows) if len(rows) >= 2 else None
